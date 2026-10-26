"""Tests for Email OSINT Lookup and Excel Export."""
import unittest
import asyncio
from app.services.email_lookup import lookup_email, export_email_intel_excel, _is_valid_email


class TestEmailLookupValidation(unittest.TestCase):
    def test_valid_emails(self):
        self.assertTrue(_is_valid_email("user@example.com"))
        self.assertTrue(_is_valid_email("first.last@company.co.uk"))
        self.assertTrue(_is_valid_email("dev+tag@sub.domain.org"))

    def test_invalid_emails(self):
        self.assertFalse(_is_valid_email("plainaddress"))
        self.assertFalse(_is_valid_email("@missingusername.com"))
        self.assertFalse(_is_valid_email("user@.com"))


class TestEmailLookupExecution(unittest.TestCase):
    def test_demo_email_lookup(self):
        data = asyncio.run(lookup_email("marcus.vance@apexnova.example"))
        self.assertEqual(data["email"], "marcus.vance@apexnova.example")
        self.assertEqual(data["domain"], "apexnova.example")
        self.assertIn("identity", data)
        self.assertEqual(data["identity"]["name"], "Marcus Vance")
        self.assertGreater(len(data["identity"]["public_webpages"]), 0)
        self.assertGreater(data["github_commits"]["total_commits"], 0)
        self.assertIn("reputation", data)
        # XposedOrNot validation
        self.assertIn("xposedornot", data)
        self.assertTrue(data["xposedornot"]["breached"])
        self.assertGreater(data["xposedornot"]["breach_count"], 0)
        self.assertIn("breaches", data)
        self.assertGreater(len(data["breaches"]), 0)

    def test_excel_export_generation(self):
        data = asyncio.run(lookup_email("marcus.vance@apexnova.example"))
        excel_bytes = export_email_intel_excel(data)
        self.assertIsInstance(excel_bytes, bytes)
        self.assertGreater(len(excel_bytes), 1000)
        # Check ZIP / XLSX magic header PK (0x50, 0x4B)
        self.assertEqual(excel_bytes[:2], b"PK")

        import io
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertIn("Identity Summary", wb.sheetnames)
        self.assertIn("Data Breaches (XON)", wb.sheetnames)
        self.assertIn("Public Webpages", wb.sheetnames)


if __name__ == "__main__":
    unittest.main()
